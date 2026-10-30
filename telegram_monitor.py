#!/usr/bin/env python3

import asyncio
import hashlib
import json
import re
import sqlite3
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

from telethon import TelegramClient, events
from telethon.tl.types import MessageEntityTextUrl

try:
    from unidecode import unidecode
except ImportError:
    unidecode = None

try:
    import pymorphy3
except ImportError:
    pymorphy3 = None

try:
    import argostranslate.package
    import argostranslate.translate
except ImportError:
    argostranslate = None



# ============================================================
# INTERFACE LANGUAGE / JEZYK INTERFEJSU
# ============================================================

LANGUAGE = "pl"


def tr(pl, en):
    return en if LANGUAGE == "en" else pl


# Automatic bilingual UI for known interface messages.
import builtins as _ui_builtins

_UI_EXACT = {
    '[ACCOUNTS WARNING] Pominięto nieprawidłowy wpis.': '[ACCOUNTS WARNING] Skipped an invalid entry.',
    '[ACCOUNTS WARNING] Pominięto konto bez nazwy.': '[ACCOUNTS WARNING] Skipped an account without a name.',
    '[WYBOR] Nieprawidlowy wybor. Podaj numery z listy, A albo 0.': '[SELECTION] Invalid choice. Enter list numbers, A, or 0.',
    '[WYBOR] Anulowano.': '[SELECTION] Cancelled.',
    '[WYBOR] Nie wybrano kont z kanalami.': '[SELECTION] No accounts with channels were selected.',
    'Rozpoczac prace z tym wyborem? [T/n]: ': 'Start with this selection? [Y/n]: ',
    'KONTO: ': 'ACCOUNT: ',
    'Konto: ': 'Account: ',
    'Wybierz kanaly konta ': 'Select channels for account ',
    ' nie ma kanalow. ': ' has no channels. ',
    'Pominieto.': 'Skipped.',
    'PODSUMOWANIE WYBORU': 'SELECTION SUMMARY',
    '[WYBOR] Brak aktywnych kont do ': '[SELECTION] No active accounts available for ',
    'Wybierz numery (np. 1,3): ': 'Select numbers (e.g. 1,3): ',
    '  A - Wszystkie': '  A - All',
    '  0 - Anuluj': '  0 - Cancel',
    'Dostepne konta Telegram:': 'Available Telegram accounts:',
    'config/accounts.json nie zawiera żadnych kont.': 'config/accounts.json contains no accounts.',
    'config/keywords.txt jest pusty.': 'config/keywords.txt is empty.',
    'Opcja --days wymaga dodatniej liczby, np. --days 3.': 'Option --days requires a positive number, e.g. --days 3.',
    'Data poczatkowa (YYYY-MM-DD): ': 'Start date (YYYY-MM-DD): ',
    'Data koncowa (YYYY-MM-DD): ': 'End date (YYYY-MM-DD): ',
    '[DATE] Data poczatkowa jest pozniejsza niz koncowa.': '[DATE] The start date is later than the end date.',
    '[DATE] Anulowano wyszukiwanie.': '[DATE] Search cancelled.',
    '[DATE] Wybierz opcje od 0 do 7.': '[DATE] Choose an option from 0 to 7.',
    '[DATE] Blad konfiguracji: ': '[DATE] Configuration error: ',
    '[DATE] Nieprawidlowa data: ': '[DATE] Invalid date: ',
    '[DATE] OD: ': '[DATE] FROM: ',
    '[DATE] DO: ': '[DATE] TO: ',
    'Łączna liczba wyników:': 'Total results:',
    'WYNIKI WEDŁUG KEYWORDÓW:': 'RESULTS BY KEYWORD:',
    'WYNIKI WEDŁUG KANAŁÓW:': 'RESULTS BY CHANNEL:',
    'STATYSTYKI': 'STATISTICS',
    'AUTOMATYCZNE GENEROWANIE:': 'AUTOMATIC GENERATION:',
    'TEST KEYWORDU:': 'KEYWORD TEST:',
    'LICZBA WARIANTÓW:': 'NUMBER OF VARIANTS:',
    'WARIANTY:': 'VARIANTS:',
    'REGEX:': 'REGEX:',
    'MONITORING AKTYWNY': 'MONITORING ACTIVE',
    'CTRL+C aby zakończyć': 'Press Ctrl+C to stop',
    'Program zatrzymany.': 'Program stopped.',
    'BŁĄD:': 'ERROR:',
    'Nieznana opcja:': 'Unknown option:',
    'Użycie: ': 'Usage: ',
}

def _ui_text(value):
    if not isinstance(value, str) or LANGUAGE != "en":
        return value
    if value in _UI_EXACT:
        return _UI_EXACT[value]
    # Translate stable prefixes in dynamic status messages.
    prefixes = (
        ("[WYBOR] Brak aktywnych kont do ", "[SELECTION] No active accounts available for "),
        ("[DATE] Blad konfiguracji: ", "[DATE] Configuration error: "),
        ("[DATE] Nieprawidlowa data: ", "[DATE] Invalid date: "),
        ("Wybierz kanaly konta ", "Select channels for account "),
        ("[INFO] Wyszukiwanie anulowane lub brak wyboru.", "[INFO] Search cancelled or nothing selected."),
        ("[INFO] Monitorowanie anulowane lub brak wyboru.", "[INFO] Monitoring cancelled or nothing selected."),
        ("Nieznana opcja:", "Unknown option:"),
        ("BŁĄD:", "ERROR:"),
    )
    for pl, en in prefixes:
        if value.startswith(pl):
            return en + value[len(pl):]
    return value

def ui_print(*args, **kwargs):
    _ui_builtins.print(*[_ui_text(arg) for arg in args], **kwargs)

def ui_input(prompt="", **kwargs):
    return _ui_builtins.input(_ui_text(prompt), **kwargs)



def choose_language():
    global LANGUAGE

    settings_path = BASE_DIR / "config" / "settings.json"
    try:
        with open(settings_path, "r", encoding="utf-8") as f:
            settings = json.load(f)
    except Exception:
        settings = {}

    saved = settings.get("interface_language", "pl")
    if saved not in ("pl", "en"):
        saved = "pl"

    ui_print()
    ui_print("=" * 50)
    ui_print("TELEGRAM MONITOR")
    ui_print("=" * 50)
    ui_print("  [1] Polski")
    ui_print("  [2] English")
    ui_print(tr(
        "  [3] Uzyj zapamietanego jezyka (" + saved + ")",
        "  [3] Use saved language (" + saved + ")"
    ))

    while True:
        choice = ui_input(
            tr("Wybierz jezyk [1/2/3]: ",
               "Choose language [1/2/3]: ")
        ).strip().lower()

        if choice == "1":
            LANGUAGE = "pl"
            break
        elif choice == "2":
            LANGUAGE = "en"
            break
        elif choice == "3":
            LANGUAGE = saved
            break
        else:
            ui_print(tr("Wpisz 1, 2 lub 3.", "Enter 1, 2, or 3."))

    settings["interface_language"] = LANGUAGE
    try:
        settings_path.parent.mkdir(parents=True, exist_ok=True)
        with open(settings_path, "w", encoding="utf-8") as f:
            json.dump(settings, f, ensure_ascii=False, indent=2)
            f.write("\n")
    except Exception as exc:
        ui_print(tr(
            "Ostrzezenie: nie zapisano jezyka: ",
            "Warning: could not save language: "
        ) + str(exc))

    ui_print(tr("Wybrany jezyk: Polski", "Selected language: English"))



BASE_DIR = Path(__file__).resolve().parent

CONFIG_DIR = BASE_DIR / "config"
RESULTS_DIR = BASE_DIR / "results"
DATABASE_DIR = BASE_DIR / "database"
CACHE_DIR = BASE_DIR / "cache"
MEDIA_DIR = BASE_DIR / "media"
VIDEOS_DIR = BASE_DIR / "videos"
SESSION_DIR = BASE_DIR / "session"

CHANNELS_FILE = CONFIG_DIR / "channels.txt"
ACCOUNTS_FILE = CONFIG_DIR / "accounts.json"
CHANNELS_DIR = CONFIG_DIR / "channels"
KEYWORDS_FILE = CONFIG_DIR / "keywords.txt"
CREDENTIALS_FILE = CONFIG_DIR / "telegram_credentials.json"
TIME_RANGE_FILE = CONFIG_DIR / "time_range.json"
SETTINGS_FILE = CONFIG_DIR / "settings.json"
TRANSLATIONS_FILE = CONFIG_DIR / "keyword_translations.json"

AUTO_KEYWORDS_FILE = CACHE_DIR / "auto_keywords.json"

DATABASE_FILE = DATABASE_DIR / "telegram_monitor.sqlite"


TARGET_LANGUAGES = [
    "pl",
    "ru",
    "uk",
    "en",
]


# ============================================================
# JSON / PLIKI
# ============================================================

def load_json(path, default=None):
    if not path.exists():
        return {} if default is None else default

    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        ui_print(f"[JSON ERROR] {path}: {e}")
        return {} if default is None else default


def save_json(path, data):
    path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    tmp = path.with_suffix(
        path.suffix + ".tmp"
    )

    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(
            data,
            f,
            ensure_ascii=False,
            indent=2
        )

    tmp.replace(path)


def load_lines(path):
    if not path.exists():
        return []

    result = []

    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()

            if not line:
                continue

            if line.startswith("#"):
                continue

            result.append(line)

    return result


def normalize_text(text):
    if not text:
        return ""

    text = text.replace("\r", "\n")
    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    return text.strip()


def sha256_text(text):
    return hashlib.sha256(
        text.encode("utf-8")
    ).hexdigest()


def safe_filename(value):
    value = re.sub(
        r'[\/:*?"<>|]+',
        "_",
        str(value)
    )

    value = value.strip()

    return value[:180] or "unknown"


def parse_date(value, end_of_day=False):
    if not value:
        return None

    value = value.strip()

    try:
        if len(value) == 10:
            dt = datetime.strptime(
                value,
                "%Y-%m-%d"
            )

            if end_of_day:
                dt = dt.replace(
                    hour=23,
                    minute=59,
                    second=59
                )

            return dt.replace(
                tzinfo=timezone.utc
            )

        dt = datetime.fromisoformat(value)

        if dt.tzinfo is None:
            dt = dt.replace(
                tzinfo=timezone.utc
            )

        return dt.astimezone(
            timezone.utc
        )

    except Exception:
        raise ValueError(
            f"Nieprawidłowa data: {value}. "
            f"Użyj YYYY-MM-DD."
        )


# ============================================================
# KONFIGURACJA
# ============================================================

def load_accounts():
    """
    Wczytuje konfigurację kont Telegram z config/accounts.json.

    Oczekiwany format:

    {
      "accounts": [
        {
          "name": "konto_1",
          "api_id": 123456,
          "api_hash": "...",
          "search_enabled": true,
          "monitor_enabled": true
        }
      ]
    }

    Kanały konta są pobierane z:
        config/channels/<name>.txt
    """

    data = load_json(
        ACCOUNTS_FILE,
        {}
    )

    if not isinstance(data, dict):
        raise RuntimeError(
            "config/accounts.json musi zawierać obiekt JSON."
        )

    accounts = data.get("accounts", [])

    if not isinstance(accounts, list):
        raise RuntimeError(
            'config/accounts.json: pole "accounts" musi być listą.'
        )

    result = []

    for account in accounts:

        if not isinstance(account, dict):
            ui_print(
                "[ACCOUNTS WARNING] Pominięto nieprawidłowy wpis."
            )
            continue

        name = str(
            account.get("name", "")
        ).strip()

        if not name:
            ui_print(
                "[ACCOUNTS WARNING] Pominięto konto bez nazwy."
            )
            continue

        item = dict(account)

        channel_file = (
            CHANNELS_DIR /
            f"{safe_filename(name)}.txt"
        )

        item["channels"] = load_lines(
            channel_file
        )

        result.append(item)

    return result



# ============================================================
# INTERAKTYWNY WYBOR KONT I KANALOW
# ============================================================

def _prompt_selection(title, items):
    while True:
        ui_print()
        ui_print(title)

        for i, item in enumerate(items, 1):
            ui_print(f"  [{i}] {item}")

        ui_print()
        ui_print(tr("  A - Wszystkie", "  A - All"))
        ui_print(tr("  0 - Anuluj", "  0 - Cancel"))

        raw = ui_input(
            tr("Wybierz numery (np. 1,3): ", "Select numbers (e.g. 1,3): ")
        ).strip().lower()

        if raw in ("a", "all"):
            return list(range(len(items)))

        if raw in ("0", "q", "quit"):
            return None

        try:
            numbers = [
                int(part.strip())
                for part in raw.split(",")
            ]

            if not numbers:
                raise ValueError

            if any(
                number < 1 or number > len(items)
                for number in numbers
            ):
                raise ValueError

            return list(dict.fromkeys(
                number - 1 for number in numbers
            ))

        except ValueError:
            ui_print(
                "[WYBOR] Nieprawidlowy wybor. "
                "Podaj numery z listy, A albo 0."
            )


def select_accounts_and_channels(accounts, mode):
    mode_label = {
        "search": "wyszukiwania",
        "monitor": "monitorowania",
    }.get(mode, mode)

    enabled_field = f"{mode}_enabled"

    eligible = [
        account
        for account in accounts
        if account.get("enabled", True)
        and account.get(enabled_field, True)
    ]

    if not eligible:
        ui_print(
            f"[WYBOR] Brak aktywnych kont do {mode_label}."
        )
        return []

    ui_print()
    ui_print("=" * 70)
    ui_print(f"WYBOR KONT DO {mode_label.upper()}")
    ui_print("=" * 70)

    account_names = [
        str(account.get("name", "bez nazwy"))
        for account in eligible
    ]

    account_indices = _prompt_selection(
        tr("Dostepne konta Telegram:", "Available Telegram accounts:"),
        account_names,
    )

    if account_indices is None:
        ui_print("[WYBOR] Anulowano.")
        return []

    selected_accounts = []

    for index in account_indices:
        account = dict(eligible[index])
        channels = list(account.get("channels") or [])
        name = str(account.get("name", "bez nazwy"))

        if not channels:
            ui_print(
                f"[WYBOR] Konto {name} nie ma kanalow. "
                "Pominieto."
            )
            continue

        ui_print()
        ui_print("=" * 70)
        ui_print(f"KONTO: {name}")
        ui_print("=" * 70)

        channel_indices = _prompt_selection(
            f"Wybierz kanaly konta {name}:",
            channels,
        )

        if channel_indices is None:
            ui_print("[WYBOR] Anulowano.")
            return []

        account["channels"] = [
            channels[i] for i in channel_indices
        ]

        selected_accounts.append(account)

    if not selected_accounts:
        ui_print("[WYBOR] Nie wybrano kont z kanalami.")
        return []

    ui_print()
    ui_print("=" * 70)
    ui_print(tr("PODSUMOWANIE WYBORU", "SELECTION SUMMARY"))
    ui_print("=" * 70)

    for account in selected_accounts:
        ui_print(f"\nKonto: {account.get('name')}")

        for channel in account["channels"]:
            ui_print(f"  - {channel}")

    ui_print()
    answer = ui_input(
        "Rozpoczac prace z tym wyborem? [T/n]: "
    ).strip().lower()

    if answer in ("n", "nie", "no", "0"):
        ui_print("[WYBOR] Anulowano.")
        return []

    return selected_accounts


def load_settings():
    data = load_json(
        SETTINGS_FILE,
        {}
    )

    if not isinstance(data, dict):
        return {}

    return data


def load_keywords():
    """
    Tylko odczytuje keywordy z config/keywords.txt.

    Aktualizacja tłumaczeń, wariantów i regexów
    jest wykonywana jawnie przez update_keyword_database().
    """

    return load_lines(KEYWORDS_FILE)


def update_keyword_database():
    """
    Jednorazowo aktualizuje bazę tłumaczeń,
    wariantów i regexów.
    """

    keywords = load_lines(KEYWORDS_FILE)

    if not keywords:
        return 0

    try:
        import auto_keywords

        ui_print("[KEYWORDS] Aktualizacja keywordów...")

        result = auto_keywords.update_all()

        if result != 0:
            ui_print(
                "[KEYWORDS] Ostrzeżenie: "
                "update_all() zwróciło kod: "
                f"{result}"
            )

        return result

    except Exception as e:
        ui_print(
            f"[KEYWORDS] Ostrzeżenie generatora: {e}"
        )
        return 1


def load_time_range():
    data = load_json(
        TIME_RANGE_FILE,
        {}
    )

    date_from = parse_date(
        data.get("from")
    )

    date_to = parse_date(
        data.get("to"),
        end_of_day=True
    )

    if not date_from or not date_to:
        raise RuntimeError(
            "config/time_range.json musi zawierać "
            '"from" oraz "to".'
        )

    if date_from > date_to:
        raise RuntimeError(
            'Data "from" jest późniejsza niż "to".'
        )

    return date_from, date_to


# ============================================================
# ARGOS
# ============================================================

def get_installed_argos_pairs():
    pairs = set()

    if argostranslate is None:
        return pairs

    try:
        packages = (
            argostranslate.package
            .get_installed_packages()
        )

        for package in packages:
            source = getattr(
                package,
                "from_code",
                None
            )

            target = getattr(
                package,
                "to_code",
                None
            )

            if source and target:
                pairs.add(
                    (source, target)
                )

    except Exception as e:
        ui_print(
            f"[ARGOS] Nie można odczytać modeli: {e}"
        )

    return pairs


def ensure_argos_models():
    if argostranslate is None:
        ui_print(
            "[ARGOS] Biblioteka nie jest dostępna."
        )
        return

    try:
        argostranslate.package.update_package_index()
    except Exception as e:
        ui_print(
            f"[ARGOS] Nie można zaktualizować indeksu: {e}"
        )
        return

    installed_pairs = (
        get_installed_argos_pairs()
    )

    available = (
        argostranslate.package
        .get_available_packages()
    )

    for source in TARGET_LANGUAGES:
        for target in TARGET_LANGUAGES:

            if source == target:
                continue

            if (source, target) in installed_pairs:
                continue

            matches = [
                p for p in available
                if getattr(
                    p,
                    "from_code",
                    None
                ) == source
                and getattr(
                    p,
                    "to_code",
                    None
                ) == target
            ]

            if not matches:
                continue

            matches.sort(
                key=lambda p: str(
                    getattr(
                        p,
                        "package_version",
                        ""
                    )
                ),
                reverse=True
            )

            package = matches[0]

            try:
                ui_print(
                    f"[ARGOS] Instaluję "
                    f"{source}->{target}..."
                )

                path = package.download()

                argostranslate.package.install_from_path(
                    path
                )

                installed_pairs.add(
                    (source, target)
                )

                ui_print(
                    f"[ARGOS] OK {source}->{target}"
                )

            except Exception as e:
                ui_print(
                    f"[ARGOS] Błąd {source}->{target}: {e}"
                )


def translate_local(
    text,
    source,
    target
):
    text = normalize_text(text)

    if not text:
        return None

    if source == target:
        return text

    if argostranslate is None:
        return None

    try:
        result = (
            argostranslate.translate.translate(
                text,
                source,
                target
            )
        )

        if result:
            return normalize_text(result)

    except Exception as e:
        ui_print(
            f"[ARGOS] "
            f"{source}->{target} "
            f"'{text}': {e}"
        )

    return None


def detect_source_language(keyword):
    """
    Prosta heurystyka.
    Keyword może być wpisany w dowolnym
    z obsługiwanych języków.
    """

    word = keyword.lower().strip()

    if re.search(
        r"[іїєґ]",
        word
    ):
        return "uk"

    if re.search(
        r"[ыэъё]",
        word
    ):
        return "ru"

    if re.search(
        r"[ąćęłńóśźż]",
        word
    ):
        return "pl"

    # Domyślnie traktujemy wpis jako polski.
    return "pl"


# ============================================================
# MORFOLOGIA
# ============================================================

def russian_forms(word):
    result = set()

    if pymorphy3 is None:
        return result

    try:
        morph = pymorphy3.MorphAnalyzer()

        parsed = morph.parse(word)

        for parse in parsed[:3]:
            normal = parse.normal_form

            if normal:
                result.add(normal)

            try:
                for grammeme_set in [
                    {"nomn", "sing"},
                    {"gent", "sing"},
                    {"datv", "sing"},
                    {"accs", "sing"},
                    {"ablt", "sing"},
                    {"loct", "sing"},
                    {"nomn", "plur"},
                    {"gent", "plur"},
                    {"datv", "plur"},
                    {"accs", "plur"},
                    {"ablt", "plur"},
                    {"loct", "plur"},
                ]:
                    inflected = parse.inflect(
                        grammeme_set
                    )

                    if inflected:
                        result.add(
                            inflected.word
                        )

            except Exception:
                pass

    except Exception as e:
        ui_print(
            f"[MORPH RU] {e}"
        )

    return result


def ukrainian_forms(word):
    result = set()

    if pymorphy3 is None:
        return result

    try:
        morph = pymorphy3.MorphAnalyzer(
            lang="uk"
        )

        parsed = morph.parse(word)

        for parse in parsed[:3]:
            normal = parse.normal_form

            if normal:
                result.add(normal)

            try:
                for grammeme_set in [
                    {"nomn", "sing"},
                    {"gent", "sing"},
                    {"datv", "sing"},
                    {"accs", "sing"},
                    {"ablt", "sing"},
                    {"loct", "sing"},
                    {"nomn", "plur"},
                    {"gent", "plur"},
                    {"datv", "plur"},
                    {"accs", "plur"},
                    {"ablt", "plur"},
                    {"loct", "plur"},
                ]:
                    inflected = parse.inflect(
                        grammeme_set
                    )

                    if inflected:
                        result.add(
                            inflected.word
                        )

            except Exception:
                pass

    except Exception as e:
        ui_print(
            f"[MORPH UK] {e}"
        )

    return result


def english_forms(word):
    result = set()

    word = word.strip().lower()

    if not word:
        return result

    # Proste, ostrożne warianty angielskie.
    if len(word) >= 3:
        if word.endswith("y"):
            result.add(
                word[:-1] + "ies"
            )
        elif word.endswith(
            ("s", "x", "z", "ch", "sh")
        ):
            result.add(
                word + "es"
            )
        else:
            result.add(
                word + "s"
            )

    return result


def polish_forms(word):
    """
    Nie udajemy pełnego analizatora języka polskiego.
    Generujemy tylko bezpieczne warianty dla prostych
    słów, aby nie produkować ogromnej liczby fałszywych
    dopasowań.
    """

    result = set()

    word = word.strip().lower()

    if not word:
        return result

    # Dla polskiego główny keyword zawsze pozostaje.
    # Dodatkowe odmiany są lepiej obsługiwane przez
    # ręczne warianty, jeśli są potrzebne.
    return result


# ============================================================
# AUTOMATYCZNE WARIANTY
# ============================================================

def load_auto_cache():
    data = load_json(
        AUTO_KEYWORDS_FILE,
        {}
    )

    if not isinstance(data, dict):
        return {}

    return data


def save_auto_cache(data):
    save_json(
        AUTO_KEYWORDS_FILE,
        data
    )


def add_variant(result, value):
    if not isinstance(
        value,
        str
    ):
        return

    value = normalize_text(
        value
    )

    if not value:
        return

    if len(value) > 100:
        return

    result.add(value)



def build_automatic_variants(keyword):
    """Wspólny generator wariantów."""
    from unified_keyword_engine import generate

    result = generate(keyword)
    return set(result["variants"])

# ============================================================
# RĘCZNE WARIANTY
# ============================================================

def flatten_variants(value):
    result = []

    if isinstance(
        value,
        str
    ):
        result.append(value)

    elif isinstance(
        value,
        list
    ):
        for item in value:

            if isinstance(
                item,
                str
            ):
                result.append(item)

    return result


def get_manual_variants(keyword):
    data = load_json(
        TRANSLATIONS_FILE,
        {}
    )

    value = data.get(
        keyword,
        {}
    )

    result = set()

    if isinstance(
        value,
        list
    ):
        for item in value:
            add_variant(
                result,
                item
            )

    elif isinstance(
        value,
        dict
    ):
        for item in value.values():

            for variant in flatten_variants(
                item
            ):
                add_variant(
                    result,
                    variant
                )

    return result


# ============================================================
# KEYWORD ENGINE
# ============================================================


def build_keyword_engine(keyword):
    """Centralny silnik dopasowywania keywordów."""
    from unified_keyword_engine import generate

    result = generate(keyword)

    return {
        "keyword": result["keyword"],
        "variants": result["variants"],
        "pattern": result["pattern"],
        "regex": result["regex"],
    }

# ============================================================
# DATABASE
# ============================================================

def init_database():
    DATABASE_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    connection = sqlite3.connect(
        DATABASE_FILE
    )

    cursor = connection.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            channel_id TEXT,
            channel_name TEXT,
            channel_username TEXT,
            message_id INTEGER,
            keyword TEXT,
            matched_variant TEXT,
            message_date TEXT,
            sender_id TEXT,
            sender_name TEXT,
            sender_username TEXT,
            post_link TEXT,
            text TEXT,
            forward_source TEXT,
            forward_message_id TEXT,
            reactions_json TEXT,
            media_type TEXT,
            media_path TEXT,
            links_json TEXT,
            context TEXT,
            summary TEXT,
            content_hash TEXT,
            created_at TEXT,
            UNIQUE(
                channel_id,
                message_id,
                keyword
            )
        )
    """)

    # Osobna tabela pilnująca wpisów w globalnym media.docx.
    # Jedna wiadomość z danego kanału może wystąpić tylko raz.
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS media_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            channel_id TEXT NOT NULL,
            message_id INTEGER NOT NULL,
            media_path TEXT,
            created_at TEXT,
            UNIQUE(
                channel_id,
                message_id
            )
        )
    """)

    connection.commit()

    return connection


def media_document_exists(
    connection,
    channel_id,
    message_id
):
    cursor = connection.cursor()

    cursor.execute("""
        SELECT 1
        FROM media_documents
        WHERE channel_id = ?
          AND message_id = ?
        LIMIT 1
    """, (
        str(channel_id),
        message_id
    ))

    return cursor.fetchone() is not None


def save_media_document_record(
    connection,
    channel_id,
    message_id,
    media_path
):
    cursor = connection.cursor()

    try:
        cursor.execute("""
            INSERT INTO media_documents (
                channel_id,
                message_id,
                media_path,
                created_at
            )
            VALUES (?, ?, ?, ?)
        """, (
            str(channel_id),
            message_id,
            media_path,
            datetime.now(
                timezone.utc
            ).isoformat()
        ))

        connection.commit()

        return True

    except sqlite3.IntegrityError:
        return False


def result_exists(
    connection,
    channel_id,
    message_id,
    keyword
):
    cursor = connection.cursor()

    cursor.execute("""
        SELECT 1
        FROM messages
        WHERE channel_id = ?
          AND message_id = ?
          AND keyword = ?
        LIMIT 1
    """, (
        str(channel_id),
        message_id,
        keyword
    ))

    return cursor.fetchone() is not None


def save_result(
    connection,
    result
):
    cursor = connection.cursor()

    try:
        cursor.execute("""
            INSERT INTO messages (
                channel_id,
                channel_name,
                channel_username,
                message_id,
                keyword,
                matched_variant,
                message_date,
                sender_id,
                sender_name,
                sender_username,
                post_link,
                text,
                forward_source,
                forward_message_id,
                reactions_json,
                media_type,
                media_path,
                links_json,
                context,
                summary,
                content_hash,
                created_at
            )
            VALUES (
                ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?, ?, ?,
                ?, ?
            )
        """, (
            str(result["channel_id"]),
            result["channel_name"],
            result["channel_username"],
            result["message_id"],
            result["keyword"],
            result["matched_variant"],
            result["message_date"],
            result["sender_id"],
            result["sender_name"],
            result["sender_username"],
            result["post_link"],
            result["text"],
            result["forward_source"],
            result["forward_message_id"],
            json.dumps(
                result["reactions"],
                ensure_ascii=False
            ),
            result["media_type"],
            result["media_path"],
            json.dumps(
                result["links"],
                ensure_ascii=False
            ),
            result["context"],
            result["summary"],
            result["content_hash"],
            datetime.now(
                timezone.utc
            ).isoformat()
        ))

        connection.commit()

        return True

    except sqlite3.IntegrityError:
        return False


# ============================================================
# TELEGRAM METADATA
# ============================================================

def get_sender_info(message):
    sender_id = None
    sender_name = None
    sender_username = None

    try:
        sender = message.sender

        if sender:

            sender_id = getattr(
                sender,
                "id",
                None
            )

            sender_username = getattr(
                sender,
                "username",
                None
            )

            first = getattr(
                sender,
                "first_name",
                ""
            )

            last = getattr(
                sender,
                "last_name",
                ""
            )

            sender_name = (
                f"{first or ''} "
                f"{last or ''}"
            ).strip()

            if not sender_name:
                sender_name = (
                    sender_username
                    or str(sender_id)
                )

    except Exception:
        pass

    return (
        sender_id,
        sender_name,
        sender_username
    )


def get_forward_info(message):
    source = None
    forward_message_id = None

    try:
        forward = message.fwd_from

        if not forward:
            return None, None

        forward_message_id = getattr(
            forward,
            "channel_post",
            None
        )

        from_id = getattr(
            forward,
            "from_id",
            None
        )

        if from_id:
            source = str(
                from_id
            )

        post_author = getattr(
            forward,
            "post_author",
            None
        )

        if post_author:
            source = (
                f"{source or ''} "
                f"{post_author}"
            ).strip()

    except Exception:
        pass

    return (
        source,
        str(forward_message_id)
        if forward_message_id is not None
        else None
    )


def get_reactions(message):
    result = []

    try:
        reactions = getattr(
            message,
            "reactions",
            None
        )

        if not reactions:
            return result

        for reaction in getattr(
            reactions,
            "results",
            []
        ):
            reaction_name = None

            reaction_obj = getattr(
                reaction,
                "reaction",
                None
            )

            if reaction_obj:

                emoticon = getattr(
                    reaction_obj,
                    "emoticon",
                    None
                )

                if emoticon:
                    reaction_name = emoticon

                else:
                    document_id = getattr(
                        reaction_obj,
                        "document_id",
                        None
                    )

                    if document_id:
                        reaction_name = (
                            f"custom:{document_id}"
                        )

            result.append({
                "reaction":
                    reaction_name or "unknown",
                "count":
                    getattr(
                        reaction,
                        "count",
                        0
                    )
            })

    except Exception:
        pass

    return result


def get_message_link(
    entity,
    message
):
    try:
        username = getattr(
            entity,
            "username",
            None
        )

        if username:
            return (
                f"https://t.me/"
                f"{username}/"
                f"{message.id}"
            )

        channel_id = getattr(
            entity,
            "id",
            None
        )

        if channel_id:
            return (
                f"https://t.me/c/"
                f"{channel_id}/"
                f"{message.id}"
            )

    except Exception:
        pass

    return None


def get_media_type(message):
    if not message.media:
        return None

    if message.photo:
        return "photo"

    if message.video:
        return "video"

    if message.document:

        mime = getattr(
            message.document,
            "mime_type",
            ""
        )

        return mime or "document"

    if message.audio:
        return "audio"

    return "media"


# ============================================================
# MEDIA
# ============================================================

async def download_media(
    message,
    channel_name,
    settings
):
    if not message.media:
        return None

    if not settings.get(
        "download_media",
        True
    ):
        return None

    media_type = get_media_type(
        message
    )

    if media_type == "video":

        max_mb = settings.get(
            "max_video_size_mb",
            1000
        )

        size = 0

        try:
            if message.document:
                size = (
                    message.document.size
                    or 0
                )
        except Exception:
            pass

        if (
            max_mb > 0
            and size > max_mb * 1024 * 1024
        ):
            ui_print(
                "[MEDIA] Pominięto film "
                "z powodu rozmiaru."
            )

            return None

    folder = (
        VIDEOS_DIR
        if media_type == "video"
        else MEDIA_DIR
    )

    folder.mkdir(
        parents=True,
        exist_ok=True
    )

    date = message.date

    filename = (
        f"{safe_filename(channel_name)}_"
        f"{date.strftime('%Y%m%d_%H%M%S')}_"
        f"msg_{message.id}"
    )

    try:

        path = await message.download_media(
            file=str(
                folder / filename
            )
        )

        if path:
            return str(
                Path(path).resolve()
            )

    except Exception as e:
        ui_print(
            f"[MEDIA ERROR] {e}"
        )

    return None


# ============================================================
# LINKI
# ============================================================

def extract_links(message):
    links = set()

    text = message.message or ""

    url_pattern = re.compile(
        r"https?://[^\s<>\"]+"
    )

    for url in url_pattern.findall(
        text
    ):
        links.add(
            url
        )

    try:

        for entity in (
            message.entities or []
        ):

            if isinstance(
                entity,
                MessageEntityTextUrl
            ):
                links.add(
                    entity.url
                )

    except Exception:
        pass

    return sorted(
        links
    )


# ============================================================
# KONTEKST
# ============================================================

async def get_context(
    client,
    entity,
    message,
    before,
    after
):
    result = []

    try:

        if before > 0:

            previous = await client.get_messages(
                entity,
                limit=before,
                offset_id=message.id
            )

            for msg in reversed(
                previous
            ):

                if msg.message:
                    result.append(
                        normalize_text(
                            msg.message
                        )
                    )

        if message.message:

            result.append(
                normalize_text(
                    message.message
                )
            )

        if after > 0:

            following = await client.get_messages(
                entity,
                limit=after,
                min_id=message.id
            )

            for msg in following:

                if msg.message:
                    result.append(
                        normalize_text(
                            msg.message
                        )
                    )

    except Exception:

        if message.message:

            result = [
                normalize_text(
                    message.message
                )
            ]

    return "\n\n".join(
        result
    )


# ============================================================
# DOCX
# ============================================================

def add_hyperlink(
    paragraph,
    text,
    url
):
    try:

        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        part = paragraph.part

        relationship_id = part.relate_to(
            url,
            "http://schemas.openxmlformats.org/"
            "officeDocument/2006/relationships/"
            "hyperlink",
            is_external=True
        )

        hyperlink = OxmlElement(
            "w:hyperlink"
        )

        hyperlink.set(
            qn("r:id"),
            relationship_id
        )

        run = OxmlElement(
            "w:r"
        )

        properties = OxmlElement(
            "w:rPr"
        )

        color = OxmlElement(
            "w:color"
        )

        color.set(
            qn("w:val"),
            "0563C1"
        )

        properties.append(
            color
        )

        underline = OxmlElement(
            "w:u"
        )

        underline.set(
            qn("w:val"),
            "single"
        )

        properties.append(
            underline
        )

        run.append(
            properties
        )

        text_node = OxmlElement(
            "w:t"
        )

        text_node.text = text

        run.append(
            text_node
        )

        hyperlink.append(
            run
        )

        paragraph._p.append(
            hyperlink
        )

    except Exception:

        paragraph.add_run(
            f"{text} ({url})"
        )


def open_document(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    if path.exists():
        return Document(
            path
        )

    document = Document()

    title = document.add_heading(
        "Telegram Monitor",
        level=1
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    document.add_paragraph(
        f"Keyword: {path.stem}"
    )

    return document


def add_metadata_line(
    document,
    label,
    value
):
    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        f"{label}: "
    )

    run.bold = True

    paragraph.add_run(
        str(value)
        if value not in (
            None,
            ""
        )
        else "brak danych"
    )


def add_result_to_doc(
    document,
    result
):
    from docx.shared import Pt, Inches

    document.add_page_break()

    document.add_heading(
        result["channel_name"],
        level=2
    )

    add_metadata_line(
        document,
        "Data",
        result["message_date"]
    )

    add_metadata_line(
        document,
        "Keyword",
        result["keyword"]
    )

    add_metadata_line(
        document,
        "Dopasowanie",
        result["matched_variant"]
    )

    add_metadata_line(
        document,
        "Autor",
        result["sender_name"]
    )

    add_metadata_line(
        document,
        "Username autora",
        result["sender_username"]
    )

    add_metadata_line(
        document,
        "ID autora",
        result["sender_id"]
    )

    if result["post_link"]:

        p = document.add_paragraph()

        p.add_run(
            "Oryginalny post: "
        ).bold = True

        add_hyperlink(
            p,
            "Otwórz w Telegramie",
            result["post_link"]
        )

    document.add_heading(
        "Forward",
        level=3
    )

    add_metadata_line(
        document,
        "Źródło",
        result["forward_source"]
    )

    add_metadata_line(
        document,
        "ID wiadomości",
        result["forward_message_id"]
    )

    document.add_heading(
        "Reakcje",
        level=3
    )

    if result["reactions"]:

        for reaction in result["reactions"]:

            document.add_paragraph(
                f"{reaction['reaction']}: "
                f"{reaction['count']}"
            )

    else:

        document.add_paragraph(
            "Brak danych."
        )

    document.add_heading(
        "Media",
        level=3
    )

    add_metadata_line(
        document,
        "Typ",
        result["media_type"]
    )

    if result["media_path"]:

        path = Path(
            result["media_path"]
        ).resolve()

        media_type = result["media_type"]

        # ====================================================
        # ZDJĘCIE
        # ====================================================

        if media_type == "photo":

            try:

                document.add_paragraph(
                    "Zdjęcie:"
                )

                paragraph = document.add_paragraph()

                run = paragraph.add_run()

                run.add_picture(
                    str(path),
                    width=Inches(5.8)
                )

                p = document.add_paragraph()

                p.add_run(
                    "Plik zdjęcia: "
                ).bold = True

                add_hyperlink(
                    p,
                    path.name,
                    "file:///"
                    + str(path).replace(
                        "\\",
                        "/"
                    )
                )

            except Exception as e:

                p = document.add_paragraph()

                p.add_run(
                    "Nie udało się osadzić zdjęcia. "
                ).bold = True

                p.add_run(
                    str(e)
                )

                p = document.add_paragraph()

                p.add_run(
                    "Plik zdjęcia: "
                ).bold = True

                add_hyperlink(
                    p,
                    path.name,
                    "file:///"
                    + str(path).replace(
                        "\\",
                        "/"
                    )
                )

        # ====================================================
        # FILM
        # ====================================================

        elif media_type == "video":

            p = document.add_paragraph()

            p.add_run(
                "Film — plik lokalny: "
            ).bold = True

            add_hyperlink(
                p,
                path.name,
                "file:///"
                + str(path).replace(
                    "\\",
                    "/"
                )
            )

        # ====================================================
        # INNE MEDIA
        # ====================================================

        else:

            p = document.add_paragraph()

            p.add_run(
                "Plik lokalny: "
            ).bold = True

            add_hyperlink(
                p,
                path.name,
                "file:///"
                + str(path).replace(
                    "\\",
                    "/"
                )
            )

    document.add_heading(
        "Linki",
        level=3
    )

    if result["links"]:

        for url in result["links"]:

            p = document.add_paragraph()

            add_hyperlink(
                p,
                url,
                url
            )

    else:

        document.add_paragraph(
            "Brak linków."
        )

    document.add_heading(
        "Kontekst",
        level=3
    )

    paragraph = document.add_paragraph(
        result["context"]
    )

    for run in paragraph.runs:
        run.font.size = Pt(10)

    document.add_heading(
        "Pełna treść",
        level=3
    )

    paragraph = document.add_paragraph(
        result["text"]
    )

    for run in paragraph.runs:
        run.font.size = Pt(10)



# ============================================================
# MEDIA — NIEZALEŻNIE OD KEYWORDÓW
# ============================================================

def open_media_document(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    if path.exists():
        return Document(path)

    document = Document()

    title = document.add_heading(
        "Telegram Monitor — MEDIA",
        level=1
    )

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(
        "Wszystkie wiadomości zawierające media "
        "z monitorowanych kanałów."
    )

    return document


async def process_media_message(
    client,
    entity,
    message,
    settings,
    document
):
    if not message.media:
        return False

    media_path = await download_media(
        message,
        getattr(
            entity,
            "title",
            str(entity)
        ),
        settings
    )

    if not media_path:
        return False

    from pathlib import Path
    from docx.shared import Pt, Inches

    message_date = message.date

    if message_date.tzinfo is None:
        message_date = message_date.replace(
            tzinfo=timezone.utc
        )

    message_date = message_date.astimezone(
        timezone.utc
    )

    (
        sender_id,
        sender_name,
        sender_username
    ) = get_sender_info(message)

    (
        forward_source,
        forward_message_id
    ) = get_forward_info(message)

    reactions = get_reactions(message)

    post_link = get_message_link(
        entity,
        message
    )

    links = extract_links(message)

    context = await get_context(
        client,
        entity,
        message,
        settings.get(
            "context_messages_before",
            1
        ),
        settings.get(
            "context_messages_after",
            1
        )
    )

    media_type = get_media_type(message)

    path = Path(media_path).resolve()

    document.add_page_break()

    document.add_heading(
        getattr(
            entity,
            "title",
            str(entity)
        ),
        level=2
    )

    add_metadata_line(
        document,
        "Data",
        message_date.strftime(
            "%Y-%m-%d %H:%M:%S"
        )
    )

    add_metadata_line(
        document,
        "ID wiadomości",
        message.id
    )

    add_metadata_line(
        document,
        "Autor",
        sender_name
    )

    add_metadata_line(
        document,
        "Username autora",
        sender_username
    )

    add_metadata_line(
        document,
        "ID autora",
        sender_id
    )

    if post_link:
        p = document.add_paragraph()
        p.add_run(
            "Oryginalny post: "
        ).bold = True

        add_hyperlink(
            p,
            "Otwórz w Telegramie",
            post_link
        )

    document.add_heading(
        "Forward",
        level=3
    )

    add_metadata_line(
        document,
        "Źródło",
        forward_source
    )

    add_metadata_line(
        document,
        "ID wiadomości",
        forward_message_id
    )

    document.add_heading(
        "Reakcje",
        level=3
    )

    if reactions:
        for reaction in reactions:
            document.add_paragraph(
                f"{reaction['reaction']}: "
                f"{reaction['count']}"
            )
    else:
        document.add_paragraph(
            "Brak danych."
        )

    document.add_heading(
        "Media",
        level=3
    )

    add_metadata_line(
        document,
        "Typ",
        media_type
    )

    if media_type == "photo":
        document.add_paragraph(
            "Zdjęcie:"
        )

        try:
            paragraph = document.add_paragraph()
            run = paragraph.add_run()

            run.add_picture(
                str(path),
                width=Inches(5.8)
            )
        except Exception as e:
            p = document.add_paragraph()
            p.add_run(
                "Nie udało się osadzić zdjęcia: "
            ).bold = True
            p.add_run(str(e))

        p = document.add_paragraph()
        p.add_run(
            "Plik zdjęcia: "
        ).bold = True

        add_hyperlink(
            p,
            path.name,
            "file:///"
            + str(path).replace(
                "\\",
                "/"
            )
        )

    else:
        p = document.add_paragraph()

        p.add_run(
            "Plik lokalny: "
        ).bold = True

        add_hyperlink(
            p,
            path.name,
            "file:///"
            + str(path).replace(
                "\\",
                "/"
            )
        )

    document.add_heading(
        "Linki",
        level=3
    )

    if links:
        for url in links:
            p = document.add_paragraph()
            add_hyperlink(
                p,
                url,
                url
            )
    else:
        document.add_paragraph(
            "Brak linków."
        )

    document.add_heading(
        "Kontekst",
        level=3
    )

    paragraph = document.add_paragraph(
        context
    )

    for run in paragraph.runs:
        run.font.size = Pt(10)

    document.add_heading(
        "Pełna treść",
        level=3
    )

    paragraph = document.add_paragraph(
        message.message or ""
    )

    for run in paragraph.runs:
        run.font.size = Pt(10)

    return str(path)


# ============================================================
# PRZETWARZANIE
# ============================================================

async def process_message(
    client,
    connection,
    entity,
    message,
    keyword,
    engine,
    settings,
    document
):
    text = normalize_text(
        message.message or ""
    )

    if not text:
        return False

    match = engine["regex"].search(
        text
    )

    if not match:
        return False

    channel_id = getattr(
        entity,
        "id",
        "unknown"
    )

    if result_exists(
        connection,
        channel_id,
        message.id,
        keyword
    ):
        return False

    channel_name = getattr(
        entity,
        "title",
        str(entity)
    )

    channel_username = getattr(
        entity,
        "username",
        None
    )

    (
        sender_id,
        sender_name,
        sender_username
    ) = get_sender_info(
        message
    )

    (
        forward_source,
        forward_message_id
    ) = get_forward_info(
        message
    )

    reactions = get_reactions(
        message
    )

    media_type = get_media_type(
        message
    )

    media_path = None

    if (
        settings.get(
            "download_media",
            True
        )
        and message.media
    ):

        media_path = await download_media(
            message,
            channel_name,
            settings
        )

    post_link = get_message_link(
        entity,
        message
    )

    links = extract_links(
        message
    )

    context = await get_context(
        client,
        entity,
        message,
        settings.get(
            "context_messages_before",
            1
        ),
        settings.get(
            "context_messages_after",
            1
        )
    )

    message_date = message.date

    if message_date.tzinfo is None:

        message_date = message_date.replace(
            tzinfo=timezone.utc
        )

    message_date = message_date.astimezone(
        timezone.utc
    )

    result = {

        "channel_id":
            channel_id,

        "channel_name":
            channel_name,

        "channel_username":
            channel_username,

        "message_id":
            message.id,

        "keyword":
            keyword,

        "matched_variant":
            match.group(0),

        "message_date":
            message_date.strftime(
                "%Y-%m-%d %H:%M:%S"
            ),

        "sender_id":
            sender_id,

        "sender_name":
            sender_name,

        "sender_username":
            sender_username,

        "post_link":
            post_link,

        "text":
            text,

        "forward_source":
            forward_source,

        "forward_message_id":
            forward_message_id,

        "reactions":
            reactions,

        "media_type":
            media_type,

        "media_path":
            media_path,

        "links":
            links,

        "context":
            context,

        "summary":
            None,

        "content_hash":
            sha256_text(
                context
            )
    }

    saved = save_result(
        connection,
        result
    )

    if saved:

        add_result_to_doc(
            document,
            result
        )

        ui_print(
            "[MATCH]",
            channel_name,
            "|",
            result["message_date"],
            "|",
            keyword,
            "|",
            match.group(0)
        )

        return True

    return False


# ============================================================
# CLIENT
# ============================================================

async def create_client(account):
    """
    Tworzy klienta Telegram dla konkretnego konta.

    Każde konto posiada własną sesję:
        session/<nazwa_konta>/telegram.session
    """

    name = str(
        account.get("name", "")
    ).strip()

    if not name:
        raise RuntimeError(
            "Konto Telegram nie ma ustawionej nazwy."
        )

    api_id = account.get(
        "api_id"
    )

    api_hash = account.get(
        "api_hash"
    )

    if not api_id:
        raise RuntimeError(
            f"Brak api_id dla konta: {name}"
        )

    if not api_hash:
        raise RuntimeError(
            f"Brak api_hash dla konta: {name}"
        )

    session_dir = (
        SESSION_DIR /
        safe_filename(name)
    )

    session_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    client = TelegramClient(
        str(
            session_dir /
            "telegram"
        ),
        int(api_id),
        api_hash
    )

    await client.start()

    return client


# ============================================================
# SEARCH
# ============================================================

async def run_search():
    update_keyword_database()

    settings = load_settings()
    accounts = load_accounts()
    accounts = select_accounts_and_channels(accounts, "search")
    if not accounts:
        ui_print(tr("[INFO] Wyszukiwanie anulowane lub brak wyboru.", "[INFO] Search cancelled or nothing selected."))
        return
    keywords = load_keywords()

    # Zakres dat do przeszukania
    days = None

    if "--days" in sys.argv:
        try:
            idx = sys.argv.index("--days")
            if idx + 1 >= len(sys.argv):
                raise ValueError
            days = int(sys.argv[idx + 1])
            if days < 1:
                raise ValueError
        except (ValueError, TypeError):
            raise RuntimeError(
                'Opcja --days wymaga dodatniej liczby, np. --days 3.'
            )

        now = datetime.now(timezone.utc)
        date_to = now
        date_from = now - timedelta(days=days)

    else:
        while True:
            ui_print()
            ui_print("=" * 70)
            ui_print(tr("WYBIERZ ZAKRES DNI DO PRZESZUKANIA", "SELECT SEARCH DATE RANGE"))
            ui_print("=" * 70)
            ui_print(tr("  [1] Ostatnie 24 godziny", "  [1] Last 24 hours"))
            ui_print(tr("  [2] Ostatnie 3 dni", "  [2] Last 3 days"))
            ui_print(tr("  [3] Ostatnie 7 dni", "  [3] Last 7 days"))
            ui_print(tr("  [4] Ostatnie 14 dni", "  [4] Last 14 days"))
            ui_print(tr("  [5] Ostatnie 30 dni", "  [5] Last 30 days"))
            ui_print(tr("  [6] Wlasny zakres dat", "  [6] Custom date range"))
            ui_print(tr("  [7] Zakres z config/time_range.json", "  [7] Use config/time_range.json"))
            ui_print(tr("  [0] Anuluj", "  [0] Cancel"))

            choice = ui_input(tr("Twoj wybor: ", "Your choice: ")).strip()

            day_options = {
                "1": 1,
                "2": 3,
                "3": 7,
                "4": 14,
                "5": 30,
            }

            if choice in day_options:
                days = day_options[choice]
                date_to = datetime.now(timezone.utc)
                date_from = date_to - timedelta(days=days)
                break

            elif choice == "6":
                try:
                    raw_from = ui_input(
                        "Data poczatkowa (YYYY-MM-DD): "
                    ).strip()
                    raw_to = ui_input(
                        "Data koncowa (YYYY-MM-DD): "
                    ).strip()

                    date_from = parse_date(raw_from)
                    date_to = parse_date(
                        raw_to,
                        end_of_day=True
                    )

                    if date_from > date_to:
                        ui_print(
                            "[DATE] Data poczatkowa jest "
                            "pozniejsza niz koncowa."
                        )
                        continue

                    break

                except ValueError as e:
                    ui_print(f"[DATE] Nieprawidlowa data: {e}")
                    continue

            elif choice == "7":
                try:
                    date_from, date_to = load_time_range()
                    break
                except (ValueError, RuntimeError) as e:
                    ui_print(f"[DATE] Blad konfiguracji: {e}")
                    continue

            elif choice == "0":
                ui_print("[DATE] Anulowano wyszukiwanie.")
                return

            else:
                ui_print("[DATE] Wybierz opcje od 0 do 7.")

    ui_print()
    ui_print("=" * 70)
    ui_print(tr("[DATE] Wybrany zakres wyszukiwania", "[DATE] Selected search range"))
    ui_print(f"[DATE] OD: {date_from.strftime('%Y-%m-%d %H:%M:%S')} UTC")
    ui_print(f"[DATE] DO: {date_to.strftime('%Y-%m-%d %H:%M:%S')} UTC")
    ui_print("=" * 70)
    ui_print()

    if not accounts:
        raise RuntimeError(
            "config/accounts.json nie zawiera żadnych kont."
        )

    if not keywords:
        raise RuntimeError(
            "config/keywords.txt jest pusty."
        )

    connection = init_database()

    clients = []

    try:

        for account in accounts:

            name = str(
                account.get("name", "")
            ).strip()

            if not account.get("enabled", True):
                ui_print(
                    f"[ACCOUNT] Pomijam konto {name}: "
                    "enabled=false"
                )
                continue

            if not account.get(
                "search_enabled",
                True
            ):
                ui_print(
                    f"[SEARCH] Pomijam konto {name}: "
                    "search_enabled=false"
                )
                continue

            channels = account.get(
                "channels",
                []
            )

            if not channels:
                ui_print(
                    f"[SEARCH WARNING] Konto {name} "
                    "nie ma przypisanych kanałów."
                )
                continue

            ui_print()
            ui_print("=" * 70)
            ui_print(
                f"[ACCOUNT] Wyszukiwanie: {name}"
            )
            ui_print("=" * 70)

            client = await create_client(
                account
            )

            clients.append(
                client
            )

            # Wspólny dokument mediów dla wszystkich kont.
            # Zdjęcia i video są zapisywane niezależnie od keywordów.
            media_result_file = (
                RESULTS_DIR /
                "media.docx"
            )

            media_document = open_media_document(
                media_result_file
            )

            engines = {}
            documents = {}

            for keyword in keywords:

                engines[keyword] = build_keyword_engine(
                    keyword
                )

                result_file = (
                    RESULTS_DIR
                    /
                    f"{safe_filename(keyword)}.docx"
                )

                documents[keyword] = open_document(
                    result_file
                )

            try:

                for channel in channels:

                    ui_print()
                    ui_print("=" * 70)
                    ui_print(
                        f"[CHANNEL] {channel}"
                    )
                    ui_print("=" * 70)

                    try:

                        entity = await client.get_entity(
                            channel
                        )

                    except Exception as e:

                        ui_print(
                            f"[ERROR] {channel}: {e}"
                        )

                        continue

                    try:

                        async for message in client.iter_messages(
                            entity
                        ):

                            if not message.date:
                                continue

                            message_date = message.date

                            if message_date.tzinfo is None:

                                message_date = message_date.replace(
                                    tzinfo=timezone.utc
                                )

                            message_date = (
                                message_date.astimezone(
                                    timezone.utc
                                )
                            )

                            if message_date < date_from:
                                break

                            if message_date > date_to:
                                continue

                            # ====================================================
                            # KEYWORDS + MEDIA
                            # ====================================================
                            # Najpierw sprawdzamy, czy wiadomość pasuje do
                            # któregokolwiek keywordu. Dzięki temu możemy
                            # zdecydować, czy media bez keywordu mają być
                            # pobierane.

                            keyword_match = False

                            for keyword in keywords:

                                try:

                                    if engines[keyword]["regex"].search(
                                        normalize_text(
                                            message.message or ""
                                        )
                                    ):
                                        keyword_match = True
                                        break

                                except Exception as e:

                                    ui_print(
                                        f"[KEYWORD CHECK ERROR] "
                                        f"{keyword}: {e}"
                                    )

                            # Media:
                            #
                            # 1. wiadomość pasuje do keywordu
                            #    -> pobierz media
                            #
                            # 2. wiadomość NIE pasuje do keywordu
                            #    -> pobierz tylko wtedy, gdy ustawienie
                            #       download_media_without_keyword=true
                            #
                            download_without_keyword = settings.get(
                                "download_media_without_keyword",
                                True
                            )

                            should_process_media = (
                                keyword_match
                                or download_without_keyword
                            )

                            if should_process_media:

                                # Globalna deduplikacja media.docx.
                                # Ta sama wiadomość z tego samego kanału
                                # może zostać zapisana tylko raz,
                                # niezależnie od konta Telegram.

                                if not media_document_exists(
                                    connection,
                                    entity.id,
                                    message.id
                                ):

                                    try:

                                        saved_media = await process_media_message(
                                            client,
                                            entity,
                                            message,
                                            settings,
                                            media_document
                                        )

                                        if saved_media:

                                            save_media_document_record(
                                                connection,
                                                entity.id,
                                                message.id,
                                                saved_media
                                            )

                                            # Zapisuj DOCX natychmiast.
                                            # Dzięki temu Ctrl+C lub przerwanie
                                            # wyszukiwania nie tracą już dodanych
                                            # wpisów z pamięci.
                                            media_document.save(
                                                media_result_file
                                            )

                                            ui_print(
                                                f"[MEDIA DOCX] "
                                                f"zapisano: {media_result_file}"
                                            )

                                    except Exception as e:

                                        ui_print(
                                            f"[MEDIA ERROR] "
                                            f"{channel}: {e}"
                                        )

                                else:

                                    ui_print(
                                        f"[MEDIA DUPLICATE] "
                                        f"{channel} | "
                                        f"message_id={message.id} "
                                        f"pominięto"
                                    )

                            # Normalne przetwarzanie keywordów.
                            # process_message() nadal posiada własną
                            # kontrolę duplikatów w SQLite.

                            for keyword in keywords:

                                try:

                                    await process_message(
                                        client,
                                        connection,
                                        entity,
                                        message,
                                        keyword,
                                        engines[keyword],
                                        settings,
                                        documents[keyword]
                                    )

                                except Exception as e:

                                    ui_print(
                                        f"[KEYWORD ERROR] "
                                        f"{keyword}: {e}"
                                    )

                    except Exception as e:

                        ui_print(
                            f"[CHANNEL ERROR] "
                            f"{channel}: {e}"
                        )

                for keyword, document in documents.items():

                    result_file = (
                        RESULTS_DIR
                        /
                        f"{safe_filename(keyword)}.docx"
                    )

                    try:

                        document.save(
                            result_file
                        )

                        ui_print(
                            f"[DOCX] {result_file}"
                        )

                    except Exception as e:

                        ui_print(
                            f"[DOCX ERROR] "
                            f"{result_file}: {e}"
                        )

                try:

                    media_document.save(
                        media_result_file
                    )

                    ui_print(
                        f"[MEDIA DOCX] "
                        f"{media_result_file}"
                    )

                except Exception as e:

                    ui_print(
                        f"[MEDIA DOCX ERROR] "
                        f"{media_result_file}: {e}"
                    )

            finally:

                await client.disconnect()

                if client in clients:
                    clients.remove(client)

    finally:

        for client in clients:

            try:
                await client.disconnect()
            except Exception:
                pass

        connection.close()


# ============================================================
# MONITOR
# ============================================================

async def run_monitor():
    update_keyword_database()

    settings = load_settings()
    accounts = load_accounts()
    accounts = select_accounts_and_channels(accounts, "monitor")
    if not accounts:
        ui_print(tr("[INFO] Monitorowanie anulowane lub brak wyboru.", "[INFO] Monitoring cancelled or nothing selected."))
        return
    keywords = load_keywords()

    if not accounts:
        raise RuntimeError(
            "config/accounts.json nie zawiera żadnych kont."
        )

    if not keywords:
        raise RuntimeError(
            "config/keywords.txt jest pusty."
        )

    connection = init_database()

    clients = []
    documents = {}

    for account in accounts:

        name = str(
            account.get("name", "")
        ).strip()

        if not account.get("enabled", True):
            ui_print(
                f"[ACCOUNT] Pomijam konto {name}: "
                "enabled=false"
            )
            continue

        if not account.get(
            "monitor_enabled",
            True
        ):
            ui_print(
                f"[MONITOR] Pomijam konto {name}: "
                "monitor_enabled=false"
            )
            continue

        channels = account.get(
            "channels",
            []
        )

        if not channels:
            ui_print(
                f"[MONITOR WARNING] Konto {name} "
                "nie ma przypisanych kanałów."
            )
            continue

        ui_print()
        ui_print("=" * 70)
        ui_print(
            f"[ACCOUNT] Monitoring: {name}"
        )
        ui_print("=" * 70)

        client = await create_client(
            account
        )

        clients.append(
            client
        )

        entities = []

        for channel in channels:

            try:

                entity = await client.get_entity(
                    channel
                )

                entities.append(
                    entity
                )

                ui_print(
                    "[MONITOR] Podłączono:",
                    channel
                )

            except Exception as e:

                ui_print(
                    "[MONITOR ERROR]",
                    channel,
                    e
                )

        if not entities:

            ui_print(
                f"[MONITOR WARNING] Konto {name}: "
                "nie udało się podłączyć do żadnego kanału."
            )

            await client.disconnect()

            if client in clients:
                clients.remove(client)

            continue

        for keyword in keywords:

            key = (
                name,
                keyword
            )

            result_file = (
                RESULTS_DIR
                /
                f"{safe_filename(keyword)}.docx"
            )

            documents[key] = open_document(
                result_file
            )

        # Zachowujemy wartości tego konta lokalnie,
        # aby każdy handler używał właściwego klienta,
        # nazwy konta i dokumentów.
        account_name = name
        account_client = client
        account_documents = documents

        @client.on(
            events.NewMessage(
                chats=entities
            )
        )
        async def handler(
            event,
            account_name=account_name,
            account_client=account_client,
            account_documents=account_documents
        ):

            message = event.message

            entity = await event.get_chat()

            for keyword in keywords:

                try:

                    key = (
                        account_name,
                        keyword
                    )

                    changed = await process_message(
                        account_client,
                        connection,
                        entity,
                        message,
                        keyword,
                        build_keyword_engine(keyword),
                        settings,
                        account_documents[key]
                    )

                    if changed:

                        result_file = (
                            RESULTS_DIR
                            /
                            f"{safe_filename(keyword)}.docx"
                        )

                        account_documents[key].save(
                            result_file
                        )

                except Exception as e:

                    ui_print(
                        f"[HANDLER ERROR] "
                        f"{account_name} / {keyword}: {e}"
                    )

    if not clients:

        connection.close()

        raise RuntimeError(
            "Nie uruchomiono żadnego konta monitoringu."
        )

    ui_print()
    ui_print("=" * 70)
    ui_print(tr("MONITORING AKTYWNY", "MONITORING ACTIVE"))
    ui_print(tr("CTRL+C aby zakończyć", "Press CTRL+C to stop"))
    ui_print("=" * 70)
    ui_print()

    try:

        await asyncio.gather(
            *[
                client.run_until_disconnected()
                for client in clients
            ]
        )

    finally:

        for (name, keyword), document in documents.items():

            result_file = (
                RESULTS_DIR
                /
                f"{safe_filename(keyword)}.docx"
            )

            try:

                document.save(
                    result_file
                )

            except Exception:
                pass

        for client in clients:

            try:
                await client.disconnect()
            except Exception:
                pass

        connection.close()


# ============================================================
# TEST KEYWORD
# ============================================================


def test_keyword(keyword):
    """
    Test keywordu.

    Przed testem automatycznie uruchamia generator,
    dzięki czemu nowe słowo dodane do keywords.txt
    dostaje tłumaczenia, warianty i regex.
    """

    keyword = keyword.strip()

    if not keyword:
        raise RuntimeError(
            "Keyword nie może być pusty."
        )

    ui_print()
    ui_print("=" * 70)
    ui_print("AUTOMATYCZNE GENEROWANIE:", keyword)
    ui_print("=" * 70)

    try:
        import auto_keywords

        result = auto_keywords.update_all()

        if result != 0:
            ui_print(
                "[WARNING] auto_keywords.py "
                f"zwrócił kod {result}"
            )

    except Exception as e:
        ui_print(
            f"[WARNING] Generator keywordów: {e}"
        )

    engine = build_keyword_engine(keyword)

    ui_print()
    ui_print("=" * 70)
    ui_print("TEST KEYWORDU:", keyword)
    ui_print("=" * 70)

    ui_print()
    ui_print("LICZBA WARIANTÓW:", len(engine["variants"]))

    ui_print()
    ui_print("WARIANTY:")

    for variant in engine["variants"]:
        ui_print(" ", variant)

    ui_print()
    ui_print("REGEX:")
    ui_print(engine["pattern"])


def rebuild_keyword_cache():
    # Najpierw ponownie generujemy tłumaczenia,
    # warianty i regexy w keyword_translations.json.
    import auto_keywords

    status = auto_keywords.update_all()

    if status != 0:
        raise RuntimeError(
            "auto_keywords.py nie zdołał odświeżyć keywordów."
        )

    keywords = load_keywords()

    if not keywords:
        raise RuntimeError(
            "config/keywords.txt jest pusty."
        )

    cache = load_auto_cache()

    for keyword in keywords:
        cache.pop(
            keyword.lower(),
            None
        )

    save_auto_cache(
        cache
    )

    ui_print(
        "[CACHE] Usunięto stare warianty."
    )

    for keyword in keywords:
        engine = build_keyword_engine(
            keyword
        )

        ui_print(
            f"[CACHE] {keyword}: "
            f"{len(engine['variants'])} wariantów"
        )


# ============================================================
# STATS
# ============================================================

def show_stats():
    connection = init_database()

    cursor = connection.cursor()

    cursor.execute(
        "SELECT COUNT(*) FROM messages"
    )

    total = cursor.fetchone()[0]

    ui_print()
    ui_print("=" * 60)
    ui_print(tr("STATYSTYKI", "STATISTICS"))
    ui_print("=" * 60)

    ui_print(
        "Łączna liczba wyników:",
        total
    )

    cursor.execute("""
        SELECT keyword, COUNT(*)
        FROM messages
        GROUP BY keyword
        ORDER BY COUNT(*) DESC
    """)

    ui_print()
    ui_print(tr("WYNIKI WEDŁUG KEYWORDÓW:", "RESULTS BY KEYWORD:"))

    for keyword, count in cursor.fetchall():

        ui_print(
            f"  {keyword}: {count}"
        )

    cursor.execute("""
        SELECT channel_name, COUNT(*)
        FROM messages
        GROUP BY channel_name
        ORDER BY COUNT(*) DESC
    """)

    ui_print()
    ui_print(tr("WYNIKI WEDŁUG KANAŁÓW:", "RESULTS BY CHANNEL:"))

    for channel, count in cursor.fetchall():

        ui_print(
            f"  {channel}: {count}"
        )

    connection.close()


# ============================================================
# HELP
# ============================================================

def print_help():
    ui_print(tr("""
TELEGRAM MONITOR V3

KOMENDY:

  python telegram_monitor.py --search
      Jednorazowe przeszukanie historii.

  python telegram_monitor.py --monitor
      Ciągłe monitorowanie nowych wiadomości.

  python telegram_monitor.py --stats
      Statystyki bazy.

  python telegram_monitor.py --test-keyword "brat"
      Test wariantów słowa kluczowego.

  python telegram_monitor.py --rebuild-cache
      Ponowne wygenerowanie wariantów.

  python telegram_monitor.py --help
      Pomoc.
""", """
TELEGRAM MONITOR V3

COMMANDS:

  python telegram_monitor.py --search
      Search message history once.

  python telegram_monitor.py --monitor
      Continuously monitor new messages.

  python telegram_monitor.py --stats
      Show database statistics.

  python telegram_monitor.py --test-keyword "brother"
      Test keyword variants.

  python telegram_monitor.py --rebuild-cache
      Regenerate keyword variants.

  python telegram_monitor.py --help
      Show help.
"""))


# ============================================================
# MAIN
# ============================================================

async def async_main():

    choose_language()

    if len(sys.argv) < 2:

        print_help()

        return

    command = sys.argv[1]

    if command == "--search":

        await run_search()

    elif command == "--monitor":

        await run_monitor()

    elif command == "--stats":

        show_stats()

    elif command == "--test-keyword":

        if len(sys.argv) < 3:

            ui_print(
                'Użycie: '
                'python telegram_monitor.py '
                '--test-keyword "brat"'
            )

            return

        test_keyword(
            sys.argv[2]
        )

    elif command == "--rebuild-cache":

        rebuild_keyword_cache()

    elif command in (
        "--help",
        "-h"
    ):

        print_help()

    else:

        ui_print(
            "Nieznana opcja:",
            command
        )

        print_help()


if __name__ == "__main__":

    try:

        asyncio.run(
            async_main()
        )

    except KeyboardInterrupt:

        ui_print(
            "\nProgram zatrzymany."
        )

    except Exception as e:

        ui_print(
            "\nBŁĄD:",
            e
        )
